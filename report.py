from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def write_report(results, cfg: dict, output_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "출결 검증결과"
    headers = ["지원자 고유값", "출결상황", "검증상태", "상세내용", "원본 파일명"]
    ws.append(headers)
    for result in results:
        attendance = cfg.get("perfect_attendance_label", "3년개근") if not any(result.values) else ", ".join(
            f"{label}{value}" for label, value in zip(cfg["attendance_labels"], result.values)
        )
        ws.append([result.applicant_id, attendance, result.status, result.detail, result.filename])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center")
    widths = [20, 110, 14, 45, 70]
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for row in ws.iter_rows(min_row=2):
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
        if row[2].value == "확인필요":
            for cell in row:
                cell.fill = PatternFill("solid", fgColor="FFF2CC")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
